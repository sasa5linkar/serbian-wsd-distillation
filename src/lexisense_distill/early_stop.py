from __future__ import annotations

import csv
import json
import shutil
import time
import traceback
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Callable, Mapping

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .evaluate import SentenceTransformerRanker, compute_metrics, evaluate_gold_split, write_errors, write_metrics, write_predictions
from .presets import LEXISENSE_PRESETS, LexiSensePreset
from .sense_repo import SenseDefinition
from .train import train_sentence_transformer


@dataclass(frozen=True, slots=True)
class EarlyStoppingJob:
    preset: LexiSensePreset
    pairs_path: Path
    gold_dir: Path
    sense_repo: Path
    best_model_dir: Path
    checkpoints_dir: Path
    outputs_dir: Path
    max_epochs: int
    patience: int
    batch_size: int
    warmup_steps: int
    text_prefix: str
    normalize_embeddings: bool
    seed: int | None

    def describe(self) -> str:
        seed_text = f" --seed {self.seed}" if self.seed is not None else ""
        return (
            f"early-stop {self.preset.name} {self.preset.model_name} "
            f"--batch-size {self.batch_size} --max-epochs {self.max_epochs} --patience {self.patience}"
            f"{seed_text}"
        )


@dataclass(frozen=True, slots=True)
class EarlyStoppingState:
    best_epoch: int = 0
    best_accuracy: float = -1.0
    epochs_without_improvement: int = 0


EpochTrainer = Callable[[EarlyStoppingJob, int, Path, Path | str], None]
ModelEvaluator = Callable[[EarlyStoppingJob, Path, str, dict[str, SenseDefinition], Path], dict[str, float | int]]


SUMMARY_COLUMNS = [
    "preset",
    "model_name",
    "best_epoch",
    "best_dev_accuracy",
    "final_accuracy",
    "final_top3_accuracy",
    "final_mrr",
    "final_coverage",
    "max_epochs",
    "patience",
    "batch_size",
    "seed",
    "train_runtime_seconds",
]


def build_early_stopping_plan(
    *,
    pairs_path: Path | str,
    gold_dir: Path | str,
    sense_repo: Path | str,
    models_dir: Path | str,
    checkpoints_dir: Path | str,
    outputs_dir: Path | str,
    max_epochs: int,
    patience: int,
    warmup_steps: int,
    seed: int | None = None,
    presets: Mapping[str, LexiSensePreset] = LEXISENSE_PRESETS,
) -> list[EarlyStoppingJob]:
    pairs_path = Path(pairs_path)
    gold_dir = Path(gold_dir)
    sense_repo = Path(sense_repo)
    models_dir = Path(models_dir)
    checkpoints_dir = Path(checkpoints_dir)
    outputs_dir = Path(outputs_dir)
    return [
        EarlyStoppingJob(
            preset=preset,
            pairs_path=pairs_path,
            gold_dir=gold_dir,
            sense_repo=sense_repo,
            best_model_dir=models_dir / f"{preset.output_name}-best",
            checkpoints_dir=checkpoints_dir / preset.name,
            outputs_dir=outputs_dir / preset.name,
            max_epochs=max_epochs,
            patience=patience,
            batch_size=preset.batch_size,
            warmup_steps=warmup_steps,
            text_prefix=preset.text_prefix,
            normalize_embeddings=preset.normalize_embeddings,
            seed=seed,
        )
        for preset in presets.values()
    ]


def update_early_stopping(
    state: EarlyStoppingState,
    *,
    epoch: int,
    dev_accuracy: float,
    min_delta: float = 0.0,
) -> EarlyStoppingState:
    if dev_accuracy > state.best_accuracy + min_delta:
        return EarlyStoppingState(best_epoch=epoch, best_accuracy=dev_accuracy, epochs_without_improvement=0)
    return EarlyStoppingState(
        best_epoch=state.best_epoch,
        best_accuracy=state.best_accuracy,
        epochs_without_improvement=state.epochs_without_improvement + 1,
    )


def should_stop(state: EarlyStoppingState, *, patience: int) -> bool:
    return state.best_epoch > 0 and state.epochs_without_improvement >= patience


def epoch_seed(base_seed: int | None, epoch: int) -> int | None:
    if base_seed is None:
        return None
    return base_seed + epoch - 1


def train_epoch_checkpoint(
    job: EarlyStoppingJob,
    epoch: int,
    checkpoint_dir: Path,
    source_model: Path | str,
) -> None:
    batch_size = job.batch_size
    try:
        train_sentence_transformer(
            job.pairs_path,
            checkpoint_dir,
            model_name=str(source_model),
            epochs=1,
            batch_size=batch_size,
            warmup_steps=job.warmup_steps,
            text_prefix=job.text_prefix,
            seed=epoch_seed(job.seed, epoch),
        )
    except RuntimeError as exc:
        if not (job.preset.name == "mling" and batch_size > 2 and "out of memory" in str(exc).lower()):
            raise
        job.outputs_dir.mkdir(parents=True, exist_ok=True)
        (job.outputs_dir / f"train_failed_epoch{epoch:02d}_batch{batch_size}.log").write_text(
            "".join(traceback.format_exception(exc)),
            encoding="utf-8",
        )
        object.__setattr__(job, "batch_size", 2)  # type: ignore[misc]
        train_sentence_transformer(
            job.pairs_path,
            checkpoint_dir,
            model_name=str(source_model),
            epochs=1,
            batch_size=2,
            warmup_steps=job.warmup_steps,
            text_prefix=job.text_prefix,
            seed=epoch_seed(job.seed, epoch),
        )


def evaluate_checkpoint(
    job: EarlyStoppingJob,
    model_dir: Path,
    split: str,
    sense_definitions: dict[str, SenseDefinition],
    out_dir: Path,
) -> dict[str, float | int]:
    ranker = SentenceTransformerRanker(
        str(model_dir),
        text_prefix=job.text_prefix,
        normalize_embeddings=job.normalize_embeddings,
    )
    predictions = evaluate_gold_split(job.gold_dir, split, sense_definitions, ranker)
    metrics = compute_metrics(predictions)
    write_metrics(metrics, out_dir / "metrics.json")
    write_predictions(predictions, out_dir / "predictions_gold.tsv")
    write_errors(predictions, out_dir / "errors_gold.tsv")
    return metrics


def copy_best_checkpoint(checkpoint_dir: Path, best_model_dir: Path) -> None:
    if best_model_dir.exists():
        shutil.rmtree(best_model_dir)
    shutil.copytree(checkpoint_dir, best_model_dir)


def run_early_stopping_job(
    job: EarlyStoppingJob,
    sense_definitions: dict[str, SenseDefinition],
    *,
    trainer: EpochTrainer = train_epoch_checkpoint,
    evaluator: ModelEvaluator = evaluate_checkpoint,
) -> dict[str, object]:
    started = time.time()
    state = EarlyStoppingState()
    epoch_rows: list[dict[str, object]] = []
    previous_model: Path | str = job.preset.model_name
    best_checkpoint: Path | None = None
    job.outputs_dir.mkdir(parents=True, exist_ok=True)

    for epoch in range(1, job.max_epochs + 1):
        checkpoint_dir = job.checkpoints_dir / f"epoch-{epoch:02d}"
        trainer(job, epoch, checkpoint_dir, previous_model)
        dev_out = job.outputs_dir / f"epoch-{epoch:02d}" / "dev_eval"
        dev_metrics = evaluator(job, checkpoint_dir, "dev", sense_definitions, dev_out)
        dev_accuracy = float(dev_metrics.get("top1_accuracy", 0.0))
        next_state = update_early_stopping(state, epoch=epoch, dev_accuracy=dev_accuracy)
        improved = next_state.best_epoch == epoch and next_state.epochs_without_improvement == 0
        if improved:
            best_checkpoint = checkpoint_dir
        state = next_state
        epoch_rows.append(
            {
                "preset": job.preset.name,
                "epoch": epoch,
                "checkpoint_dir": str(checkpoint_dir),
                "dev_top1_accuracy": dev_accuracy,
                "dev_top3_accuracy": dev_metrics.get("top3_accuracy", 0.0),
                "dev_mrr": dev_metrics.get("mrr", 0.0),
                "dev_coverage": dev_metrics.get("coverage", 0.0),
                "seed": job.seed,
                "epoch_seed": epoch_seed(job.seed, epoch),
                "is_best": improved,
                "epochs_without_improvement": state.epochs_without_improvement,
            }
        )
        previous_model = checkpoint_dir
        if should_stop(state, patience=job.patience):
            break

    if best_checkpoint is None:
        raise RuntimeError(f"No best checkpoint selected for {job.preset.name}")
    copy_best_checkpoint(best_checkpoint, job.best_model_dir)
    final_metrics = evaluator(job, job.best_model_dir, "final", sense_definitions, job.outputs_dir / "final_eval")
    best_dev_metrics = evaluator(job, job.best_model_dir, "dev", sense_definitions, job.outputs_dir / "dev_eval")
    runtime = round(time.time() - started, 3)

    write_epoch_metrics(epoch_rows, job.outputs_dir / "epoch_metrics.tsv")
    train_summary = {
        "preset": job.preset.name,
        "model_name": job.preset.model_name,
        "best_model_dir": str(job.best_model_dir),
        "best_epoch": state.best_epoch,
        "best_dev_accuracy": state.best_accuracy,
        "epochs_ran": len(epoch_rows),
        "max_epochs": job.max_epochs,
        "patience": job.patience,
        "batch_size": job.batch_size,
        "warmup_steps": job.warmup_steps,
        "text_prefix": job.text_prefix,
        "normalize_embeddings": job.normalize_embeddings,
        "seed": job.seed,
        "train_runtime_seconds": runtime,
        "best_dev_metrics": best_dev_metrics,
        "final_metrics": final_metrics,
    }
    (job.outputs_dir / "train_summary.json").write_text(
        json.dumps(train_summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return train_summary


def write_epoch_metrics(rows: list[dict[str, object]], path: Path | str) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "preset",
        "epoch",
        "checkpoint_dir",
        "dev_top1_accuracy",
        "dev_top3_accuracy",
        "dev_mrr",
        "dev_coverage",
        "seed",
        "epoch_seed",
        "is_best",
        "epochs_without_improvement",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, delimiter="\t", fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def collect_early_stopping_summary(outputs_dir: Path | str, presets: Mapping[str, LexiSensePreset] = LEXISENSE_PRESETS) -> list[dict[str, object]]:
    outputs_dir = Path(outputs_dir)
    rows: list[dict[str, object]] = []
    for preset in presets.values():
        path = outputs_dir / preset.name / "train_summary.json"
        if not path.exists():
            continue
        summary = json.loads(path.read_text(encoding="utf-8"))
        final_metrics = summary.get("final_metrics", {})
        rows.append(
            {
                "preset": preset.name,
                "model_name": preset.model_name,
                "best_epoch": summary.get("best_epoch", 0),
                "best_dev_accuracy": summary.get("best_dev_accuracy", 0.0),
                "final_accuracy": final_metrics.get("top1_accuracy", 0.0),
                "final_top3_accuracy": final_metrics.get("top3_accuracy", 0.0),
                "final_mrr": final_metrics.get("mrr", 0.0),
                "final_coverage": final_metrics.get("coverage", 0.0),
                "max_epochs": summary.get("max_epochs", 0),
                "patience": summary.get("patience", 0),
                "batch_size": summary.get("batch_size", 0),
                "seed": summary.get("seed", ""),
                "train_runtime_seconds": summary.get("train_runtime_seconds", 0.0),
            }
        )
    return rows


def write_early_stopping_summary(rows: list[dict[str, object]], tsv_path: Path | str, json_path: Path | str) -> None:
    tsv_path = Path(tsv_path)
    json_path = Path(json_path)
    tsv_path.parent.mkdir(parents=True, exist_ok=True)
    with tsv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, delimiter="\t", fieldnames=SUMMARY_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)
    json_path.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")


def _read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def collect_epoch_metric_rows(outputs_dir: Path | str) -> list[dict[str, object]]:
    outputs_dir = Path(outputs_dir)
    rows: list[dict[str, object]] = []
    for path in sorted(outputs_dir.glob("*/epoch_metrics.tsv")):
        rows.extend(_read_tsv(path))
    return rows


def _write_workbook(path: Path, sheets: dict[str, list[dict[str, object]]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    percent_fields = {
        "best_dev_accuracy",
        "final_accuracy",
        "final_top3_accuracy",
        "final_mrr",
        "final_coverage",
        "dev_top1_accuracy",
        "dev_top3_accuracy",
        "dev_mrr",
        "dev_coverage",
    }
    for sheet_name, rows in sheets.items():
        worksheet = workbook.create_sheet(sheet_name)
        fieldnames: list[str] = []
        for row in rows:
            for key in row:
                if key not in fieldnames:
                    fieldnames.append(key)
        if not fieldnames:
            fieldnames = ["note"]
            rows = [{"note": "No rows"}]
        worksheet.append(fieldnames)
        for row in rows:
            worksheet.append([row.get(field, "") for field in fieldnames])
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for column_index, field in enumerate(fieldnames, start=1):
            width = min(60, max(12, len(field) + 2, *(len(str(row.get(field, ""))) + 2 for row in rows[:200])))
            worksheet.column_dimensions[get_column_letter(column_index)].width = width
            if field in percent_fields:
                for cells in worksheet.iter_cols(min_col=column_index, max_col=column_index, min_row=2):
                    for cell in cells:
                        cell.number_format = "0.0%"
    workbook.save(path)


def write_early_stopping_workbooks(outputs_dir: Path | str, summary_rows: list[dict[str, object]]) -> None:
    outputs_dir = Path(outputs_dir)
    epoch_rows = collect_epoch_metric_rows(outputs_dir)
    _write_workbook(outputs_dir / "summary.xlsx", {"Summary": summary_rows, "Training Curves": epoch_rows})
    _write_workbook(outputs_dir / "training_curves.xlsx", {"Training Curves": epoch_rows})


def plan_to_jsonable(plan: list[EarlyStoppingJob]) -> list[dict[str, object]]:
    return [
        {
            "preset": asdict(job.preset),
            "pairs_path": str(job.pairs_path),
            "gold_dir": str(job.gold_dir),
            "sense_repo": str(job.sense_repo),
            "best_model_dir": str(job.best_model_dir),
            "checkpoints_dir": str(job.checkpoints_dir),
            "outputs_dir": str(job.outputs_dir),
            "max_epochs": job.max_epochs,
            "patience": job.patience,
            "batch_size": job.batch_size,
            "warmup_steps": job.warmup_steps,
            "seed": job.seed,
            "text_prefix": job.text_prefix,
            "normalize_embeddings": job.normalize_embeddings,
            "description": job.describe(),
        }
        for job in plan
    ]

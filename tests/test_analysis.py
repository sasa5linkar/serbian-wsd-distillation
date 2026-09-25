from __future__ import annotations

import csv
import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from lexisense_distill.analysis import (
    EXCEL_SHEETS,
    align_predictions_with_gold,
    compute_stratified_metrics,
    holm_adjust,
    mcnemar_test,
    sense_bucket,
    write_analysis_outputs,
    write_excel_workbook,
    wilson_interval,
)
from lexisense_distill.evaluate import evaluate_raw_round2_workbook
from lexisense_distill.sense_repo import SenseDefinition


def _write_gold(path: Path) -> None:
    path.write_text(
        "\n".join(
            [
                "#FORMAT=WebAnno TSV 3.3",
                "#T_SP=webanno.custom.WSD|Comment|Explanation|KBid|NumberOfSenses|Origine|Possible",
                "#Text=Prva rečenica ima duplu metu.",
                "1-1\t0-4\tPrva\tX\tADJ\t_\t_\tprvi\t_\t_\t_\t_\t_\tENG30-00000001-a\t2\t_\tENG30-00000001-a;ENG30-00000002-a",
                "1-2\t5-13\trečenica\tX\tNOUN\t_\t_\trečenica\t_\t_\t_\t_\t_\tENG30-00000003-n\t_\t_\tENG30-00000003-n;ENG30-00000004-n;ENG30-00000005-n",
                "1-3\t14-17\tima\tX\tVERB\t_\t_\timati\t_\t_\t_\t_\t_\t_\t_\t_\t_",
                "1-4\t18-23\tduplu\tX\tADJ\t_\t_\tdupli\t_\t_\t_\t_\t_\t_\t_\t_\t_",
                "1-5\t24-28\tmetu\tX\tNOUN\t_\t_\tmeta\t_\t_\t_\t_\t_\tENG30-00000003-n\t3\t_\tENG30-00000003-n;ENG30-00000004-n;ENG30-00000005-n",
                "",
                "#Text=Multi word primer.",
                "2-1\t0-5\tMulti\tX\tADJ\t_\t_\tmulti\t1\tmulti word\tMWE\t_\t_\t_\t_\t_\t_",
                "2-2\t6-10\tword\tX\tNOUN\t_\t_\tword\t1\tmulti word\tMWE\t_\t_\tLLM-0001\t4\t_\tLLM-0001;LLM-0002;LLM-0003;LLM-0004",
                "2-3\t11-17\tprimer\tX\tNOUN\t_\t_\tprimer\t_\t_\t_\t_\t_\t_\t_\t_\t_",
            ]
        ),
        encoding="utf-8",
    )


def _write_predictions(path: Path) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            delimiter="\t",
            fieldnames=["sentence", "target", "gold_sense_id", "ranked_sense_ids", "covered"],
        )
        writer.writeheader()
        writer.writerows(
            [
                {
                    "sentence": "Prva rečenica ima duplu metu.",
                    "target": "Prva",
                    "gold_sense_id": "ENG30-00000001-a",
                    "ranked_sense_ids": "ENG30-00000001-a;ENG30-00000002-a",
                    "covered": "True",
                },
                {
                    "sentence": "Prva rečenica ima duplu metu.",
                    "target": "rečenica",
                    "gold_sense_id": "ENG30-00000003-n",
                    "ranked_sense_ids": "ENG30-00000004-n;ENG30-00000003-n",
                    "covered": "True",
                },
                {
                    "sentence": "Prva rečenica ima duplu metu.",
                    "target": "metu",
                    "gold_sense_id": "ENG30-00000003-n",
                    "ranked_sense_ids": "ENG30-00000005-n;ENG30-00000004-n;ENG30-00000003-n",
                    "covered": "True",
                },
                {
                    "sentence": "Multi word primer.",
                    "target": "multi word",
                    "gold_sense_id": "LLM-0001",
                    "ranked_sense_ids": "",
                    "covered": "False",
                },
            ]
        )


def test_align_predictions_enriches_duplicate_targets_sense_counts_mwe_and_upos(tmp_path: Path) -> None:
    gold = tmp_path / "gold.tsv"
    predictions = tmp_path / "predictions.tsv"
    _write_gold(gold)
    _write_predictions(predictions)

    rows = align_predictions_with_gold(
        prediction_path=predictions,
        gold_paths=[gold],
        split="dev",
        model="simple",
        model_name="all-MiniLM-L6-v2",
    )

    assert [row.example_id for row in rows] == ["dev-000001", "dev-000002", "dev-000003", "dev-000004"]
    assert rows[0].number_of_senses == 2
    assert rows[1].number_of_senses == 3
    assert rows[2].number_of_senses == 3
    assert rows[3].number_of_senses == 4
    assert rows[3].target_type == "mwe"
    assert rows[3].upos == "ADJ+NOUN"
    assert rows[3].lemma == "multi word"
    assert rows[1].top1_correct is False
    assert rows[2].top3_correct is True
    assert rows[3].covered is False


def test_sense_bucket_groups_ambiguity_levels() -> None:
    assert sense_bucket(1) == "1"
    assert sense_bucket(2) == "2"
    assert sense_bucket(3) == "3"
    assert sense_bucket(5) == "4-5"
    assert sense_bucket(10) == "6-10"
    assert sense_bucket(11) == "11+"


def test_compute_stratified_metrics_reports_top1_counts_and_ranking_metrics(tmp_path: Path) -> None:
    gold = tmp_path / "gold.tsv"
    predictions = tmp_path / "predictions.tsv"
    _write_gold(gold)
    _write_predictions(predictions)
    rows = align_predictions_with_gold(predictions, [gold], "dev", "simple", "all-MiniLM-L6-v2")

    metrics = compute_stratified_metrics(rows, ["split", "model"])

    assert metrics == [
        {
            "split": "dev",
            "model": "simple",
            "total_examples": 4,
            "covered_examples": 3,
            "coverage": 0.75,
            "correct_top1": 1,
            "incorrect_top1": 2,
            "accuracy": 1 / 3,
            "accuracy_ci_low": wilson_interval(1, 3)[0],
            "accuracy_ci_high": wilson_interval(1, 3)[1],
            "accuracy_ci_95": "6.1-79.2",
            "top3_accuracy": 1.0,
            "mrr": (1.0 + 0.5 + 1 / 3) / 3,
        }
    ]


def test_wilson_interval_returns_95_percent_bounds() -> None:
    low, high = wilson_interval(1079, 1921)

    assert round(low * 100, 1) == 53.9
    assert round(high * 100, 1) == 58.4


def test_mcnemar_and_holm_adjustment() -> None:
    a = [True, True, False, False, True, False]
    b = [True, False, True, True, False, False]

    result = mcnemar_test(a, b, model_a="a", model_b="b", split="final")

    assert result["n01"] == 2
    assert result["n10"] == 2
    assert result["discordant"] == 4
    assert result["chi_square_cc"] == 0.25
    assert result["exact_p_value"] == 1.0
    assert result["accuracy_diff"] == 0.0

    adjusted = holm_adjust(
        [
            {"split": "final", "p_value": 0.01},
            {"split": "final", "p_value": 0.04},
            {"split": "final", "p_value": 0.03},
            {"split": "dev", "p_value": 0.02},
        ]
    )
    final_rows = [row for row in adjusted if row["split"] == "final"]
    assert [row["holm_p_value"] for row in final_rows] == [0.03, 0.06, 0.06]


def test_write_excel_workbook_creates_required_sheets(tmp_path: Path) -> None:
    workbook = tmp_path / "analysis.xlsx"
    tables = {sheet: [{"split": "dev", "model": "simple", "accuracy": 0.5}] for sheet in EXCEL_SHEETS}

    write_excel_workbook(tables, workbook)

    book = load_workbook(workbook)
    assert book.sheetnames == list(EXCEL_SHEETS)
    assert book["Summary"].freeze_panes == "A2"
    assert book["Summary"].auto_filter.ref == "A1:C2"


def test_write_analysis_outputs_smoke(tmp_path: Path) -> None:
    comparison = tmp_path / "comparison"
    gold_dir = tmp_path / "gold"
    out_dir = tmp_path / "analysis"
    gold_dir.mkdir()
    (gold_dir / "sr-elexis-WSD_0001_0120-gold.tsv").write_text("", encoding="utf-8")
    (gold_dir / "sr-elexis-WSD_0301-0400-gold.tsv").write_text("", encoding="utf-8")
    (gold_dir / "sr-elexis-WSD_0401-0500-gold.tsv").write_text("", encoding="utf-8")
    (gold_dir / "sr-elexis-WSD_0501-0600-gold.tsv").write_text("", encoding="utf-8")
    _write_gold(gold_dir / "sr-elexis-WSD_0001_0120-gold.tsv")
    for split in ["dev", "final"]:
        for preset in ["simple", "tesla", "mling"]:
            eval_dir = comparison / preset / f"{split}_eval"
            eval_dir.mkdir(parents=True)
            _write_predictions(eval_dir / "predictions_gold.tsv")
    (comparison / "summary.json").write_text(
        json.dumps(
            [
                {"preset": "simple", "model_name": "all-MiniLM-L6-v2"},
                {"preset": "tesla", "model_name": "te-sla/TeslaXLM"},
                {"preset": "mling", "model_name": "intfloat/multilingual-e5-large"},
            ]
        ),
        encoding="utf-8",
    )

    summary = write_analysis_outputs(comparison, gold_dir, out_dir)

    assert summary["models"] == ["mling", "simple", "tesla"]
    assert (out_dir / "summary_metrics.tsv").exists()
    assert (out_dir / "stratified_metrics.tsv").exists()
    assert (out_dir / "mcnemar_tests.tsv").exists()
    assert (out_dir / "errors_enriched.tsv").exists()
    assert (out_dir / "analysis_summary.json").exists()
    assert (out_dir / "wsd_deep_eval.xlsx").exists()
    assert (out_dir / "training_methods_report.md").exists()
    assert (out_dir / "raw_round2_baselines.tsv").exists()
    assert (out_dir / "raw_round2_by_sense_count.tsv").exists()
    assert (out_dir / "raw_round2_by_sense_bucket.tsv").exists()
    assert (out_dir / "before_after_distillation.tsv").exists()


def test_write_analysis_outputs_can_use_raw_round2_workbook(tmp_path: Path) -> None:
    comparison = tmp_path / "comparison"
    gold_dir = tmp_path / "gold"
    out_dir = tmp_path / "analysis"
    gold_dir.mkdir()
    _write_gold(gold_dir / "sr-elexis-WSD_0001_0120-gold.tsv")
    for name in ["sr-elexis-WSD_0301-0400-gold.tsv", "sr-elexis-WSD_0401-0500-gold.tsv", "sr-elexis-WSD_0501-0600-gold.tsv"]:
        (gold_dir / name).write_text("", encoding="utf-8")
    for split in ["dev", "final"]:
        for preset in ["simple", "tesla", "mling"]:
            eval_dir = comparison / preset / f"{split}_eval"
            eval_dir.mkdir(parents=True)
            _write_predictions(eval_dir / "predictions_gold.tsv")
    (comparison / "summary.json").write_text(
        json.dumps(
            [
                {"preset": "simple", "model_name": "all-MiniLM-L6-v2"},
                {"preset": "tesla", "model_name": "te-sla/TeslaXLM"},
                {"preset": "mling", "model_name": "intfloat/multilingual-e5-large"},
            ]
        ),
        encoding="utf-8",
    )
    raw_xlsx = tmp_path / "raw_round2.xlsx"
    pd.DataFrame(
        [
            {
                "sent_id": 301,
                "span": "0-4",
                "form": "Dobar",
                "Xpos": "A",
                "Upos": "ADJ",
                "Lemma": "dobar",
                "type": "s",
                "MWE_type": "",
                "NoS_rep2": 2,
                "NER": "_",
                "Gold_Sense": "seENG3000000001a",
                "Gold_Possible": "ENG30-00000001-a;ENG30-00000002-a",
                "Sense_simple_wsd": "seENG3000000001a",
                "Sense_tesla_wsd": "seENG3000000002a",
                "Sense_mling": "seENG3000000001a",
            },
            {
                "sent_id": 301,
                "span": "5-10",
                "form": "Beograd",
                "Xpos": "N",
                "Upos": "NOUN",
                "Lemma": "Beograd",
                "type": "s",
                "MWE_type": "",
                "NoS_rep2": 3,
                "NER": "LOC",
                "Gold_Sense": "seENG3000000003n",
                "Gold_Possible": "ENG30-00000003-n;ENG30-00000004-n;ENG30-00000005-n",
                "Sense_simple_wsd": "seENG3000000003n",
                "Sense_tesla_wsd": "seENG3000000003n",
                "Sense_mling": "seENG3000000003n",
            },
        ]
    ).to_excel(raw_xlsx, sheet_name="All_models_wide", index=False)

    summary = write_analysis_outputs(comparison, gold_dir, out_dir, raw_round2_xlsx=raw_xlsx)

    simple_raw = [
        row
        for row in summary["raw_round2_baselines"]
        if row["model"] == "simple" and row["target_type"] == "total"
    ][0]
    assert simple_raw["matched"] == 1
    assert simple_raw["total_matched_mismatched"] == 1
    assert summary["raw_round2_by_sense_count"][0]["number_of_senses"] == 2


def test_evaluate_raw_round2_workbook_uses_content_blank_ner_universe(tmp_path: Path) -> None:
    gold_dir = tmp_path / "gold"
    gold_dir.mkdir()
    (gold_dir / "sr-elexis-WSD_0001_0120-gold.tsv").write_text("", encoding="utf-8")
    (gold_dir / "sr-elexis-WSD_0401-0500-gold.tsv").write_text("", encoding="utf-8")
    (gold_dir / "sr-elexis-WSD_0501-0600-gold.tsv").write_text("", encoding="utf-8")
    (gold_dir / "sr-elexis-WSD_0301-0400-gold.tsv").write_text(
        "\n".join(
            [
                "#FORMAT=WebAnno TSV 3.3",
                "#T_SP=webanno.custom.WSD|Comment|Explanation|KBid|NumberOfSenses|Origine|Possible",
                "#Text=Dobar test Beograd.",
                "1-1\t0-5\tDobar\tA\tADJ\t_\t_\tdobar\t_\t_\t_\t_\t_\tENG30-00000001-a\t2\t_\tENG30-00000001-a;ENG30-00000002-a",
                "1-2\t6-10\ttest\tN\tNOUN\t_\t_\ttest\t_\t_\t_\t_\t_\tENG30-00000003-n\t2\t_\tENG30-00000003-n;ENG30-00000004-n",
                "1-3\t11-18\tBeograd\tN\tNOUN\t_\tLOC\tBeograd\t_\t_\t_\t_\t_\tENG30-00000005-n\t1\t_\tENG30-00000005-n",
            ]
        ),
        encoding="utf-8",
    )
    raw_xlsx = tmp_path / "raw.xlsx"
    pd.DataFrame(
        [
            {
                "sent_id": 301,
                "form": "Dobar",
                "Upos": "ADJ",
                "type": "s",
                "NER": "_",
                "Gold_Sense": "ENG30-00000001-a",
                "Gold_Possible": "ENG30-00000001-a;ENG30-00000002-a",
                "norm_span": "0-5",
            },
            {
                "sent_id": 301,
                "form": "Beograd",
                "Upos": "NOUN",
                "type": "s",
                "NER": "LOC",
                "Gold_Sense": "ENG30-00000005-n",
                "Gold_Possible": "ENG30-00000005-n",
                "norm_span": "11-18",
            },
            {
                "sent_id": 301,
                "form": "test",
                "Upos": "NOUN",
                "type": "s",
                "NER": "_",
                "Gold_Sense": "",
                "Gold_Possible": "",
                "norm_span": "6-10",
            },
        ]
    ).to_excel(raw_xlsx, sheet_name="All_models_wide", index=False)

    class FirstRanker:
        def __init__(self) -> None:
            self.texts: list[str] = []

        def rank(self, text: str, definitions: list[str]) -> list[int]:
            self.texts.append(text)
            return list(range(len(definitions)))

    ranker = FirstRanker()
    predictions = evaluate_raw_round2_workbook(
        raw_xlsx,
        gold_dir,
        {
            "ENG30-00000001-a": SenseDefinition("ENG30-00000001-a", "dobar", "ADJ", "good"),
            "ENG30-00000002-a": SenseDefinition("ENG30-00000002-a", "dobar", "ADJ", "kind"),
        },
        ranker,
    )

    assert len(predictions) == 1
    assert predictions[0].target == "Dobar"
    assert ranker.texts == ["**Dobar** test Beograd."]


def test_evaluate_raw_round2_workbook_short_circuits_single_candidate_without_definition(tmp_path: Path) -> None:
    gold_dir = tmp_path / "gold"
    gold_dir.mkdir()
    (gold_dir / "sr-elexis-WSD_0001_0120-gold.tsv").write_text("", encoding="utf-8")
    (gold_dir / "sr-elexis-WSD_0401-0500-gold.tsv").write_text("", encoding="utf-8")
    (gold_dir / "sr-elexis-WSD_0501-0600-gold.tsv").write_text("", encoding="utf-8")
    (gold_dir / "sr-elexis-WSD_0301-0400-gold.tsv").write_text(
        "\n".join(
            [
                "#FORMAT=WebAnno TSV 3.3",
                "#T_SP=webanno.custom.WSD|Comment|Explanation|KBid|NumberOfSenses|Origine|Possible",
                "#Text=Jedan kandidat.",
                "1-1\t0-5\tJedan\tA\tADJ\t_\t_\tjedan\t_\t_\t_\t_\t_\tENG30-00000001-a\t1\t_\tENG30-00000001-a",
            ]
        ),
        encoding="utf-8",
    )
    raw_xlsx = tmp_path / "raw.xlsx"
    pd.DataFrame(
        [
            {
                "sent_id": 301,
                "form": "Jedan",
                "Upos": "ADJ",
                "type": "s",
                "NER": "_",
                "Gold_Sense": "ENG30-00000001-a",
                "Gold_Possible": "ENG30-00000001-a",
                "norm_span": "0-5",
            }
        ]
    ).to_excel(raw_xlsx, sheet_name="All_models_wide", index=False)

    class UnusedRanker:
        def rank(self, text: str, definitions: list[str]) -> list[int]:
            raise AssertionError("single-candidate examples must not be ranked")

    predictions = evaluate_raw_round2_workbook(raw_xlsx, gold_dir, {}, UnusedRanker())

    assert predictions[0].covered is True
    assert predictions[0].ranked_sense_ids == ["ENG30-00000001-a"]
